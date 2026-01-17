import os
import pytest
from docx import Document
from academic_proofreader.word_processor import WordProcessor

def test_read_paragraphs(tmp_path):
    """Test reading text from paragraphs in a Word document."""
    doc_path = tmp_path / "test.docx"
    doc = Document()
    doc.add_paragraph("Paragraph 1")
    doc.add_paragraph("Paragraph 2")
    doc.save(doc_path)
    
    processor = WordProcessor(doc_path)
    text = processor.get_text()
    
    assert "Paragraph 1" in text
    assert "Paragraph 2" in text
    assert text.strip() == "Paragraph 1\nParagraph 2"

def test_save_document(tmp_path):
    """Test saving the document to a new file."""
    doc_path = tmp_path / "test.docx"
    doc = Document()
    doc.add_paragraph("Original Text")
    doc.save(doc_path)
    
    processor = WordProcessor(doc_path)
    save_path = tmp_path / "saved.docx"
    processor.save(save_path)
    
    assert os.path.exists(save_path)
    new_doc = Document(save_path)
    assert new_doc.paragraphs[0].text == "Original Text"



def test_add_tracked_change(tmp_path):
    """Test adding a tracked change (revision) to a document."""
    from lxml import etree
    
    doc_path = tmp_path / "test_tracked.docx"
    doc = Document()
    doc.add_paragraph("This is the original text that needs correction.")
    doc.save(doc_path)
    
    processor = WordProcessor(doc_path)
    result = processor.add_tracked_change("original text", "corrected text", author="Test Author")
    
    assert result is True
    
    save_path = tmp_path / "tracked.docx"
    processor.save(save_path)
    
    assert os.path.exists(save_path)
    
    # Verify the document by checking for w:del and w:ins elements in the XML
    tracked_doc = Document(save_path)
    para = tracked_doc.paragraphs[0]
    para_xml = para._p.xml
    
    # Check for deletion and insertion markers
    assert 'w:del' in para_xml
    assert 'w:ins' in para_xml
    assert 'w:delText' in para_xml
    assert 'original text' in para_xml
    assert 'corrected text' in para_xml


def test_add_tracked_change_not_found(tmp_path):
    """Test that add_tracked_change returns False when text is not found."""
    doc_path = tmp_path / "test_not_found.docx"
    doc = Document()
    doc.add_paragraph("Some other text here.")
    doc.save(doc_path)
    
    processor = WordProcessor(doc_path)
    result = processor.add_tracked_change("nonexistent text", "replacement", author="Test Author")
    
    assert result is False


def test_enable_track_revisions(tmp_path):
    """Test enabling track revisions in document settings."""
    doc_path = tmp_path / "test_revisions.docx"
    doc = Document()
    doc.add_paragraph("Test paragraph.")
    doc.save(doc_path)
    
    processor = WordProcessor(doc_path)
    result = processor.enable_track_revisions()
    
    assert result is True
    
    save_path = tmp_path / "revisions_enabled.docx"
    processor.save(save_path)
    
    # Verify by checking settings.xml contains trackRevisions
    from zipfile import ZipFile
    with ZipFile(save_path, 'r') as docx:
        settings_xml = docx.read('word/settings.xml').decode('utf-8')
        assert 'trackRevisions' in settings_xml
