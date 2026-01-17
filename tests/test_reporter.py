import pytest
from docx import Document
from docx.document import Document as DocumentClass
from academic_proofreader.reporter import generate_report, save_report

def test_generate_report():
    """Test generating a Word summary report."""
    summary = {
        "overall_assessment": "The document is well-written.",
        "strengths": ["Clear structure", "Good vocabulary"],
        "areas_for_improvement": ["Minor grammar issues"]
    }
    issues = [
        {"category": "Grammar", "explanation": "E1"},
        {"category": "Style", "explanation": "E2"},
        {"category": "Grammar", "explanation": "E3"}
    ]
    
    report = generate_report(summary, issues)
    
    # Should return a Document object
    assert isinstance(report, DocumentClass)
    
    # Extract all text from the document
    all_text = "\n".join([p.text for p in report.paragraphs])
    
    assert "well-written" in all_text
    assert "Clear structure" in all_text
    assert "Minor grammar issues" in all_text
    assert "Total corrections applied: 3" in all_text
    assert "Grammar: 2" in all_text
    assert "Style: 1" in all_text


def test_save_report(tmp_path):
    """Test saving a report to a file."""
    summary = {"overall_assessment": "Good."}
    issues = []
    
    report = generate_report(summary, issues)
    output_path = tmp_path / "test_report.docx"
    save_report(report, output_path)
    
    assert output_path.exists()
    
    # Verify it can be opened
    loaded = Document(output_path)
    assert len(loaded.paragraphs) > 0


