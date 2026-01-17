from collections import Counter
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT


def generate_report(summary, issues):
    """
    Generate a high-level Word document summary report of the proofreading findings.
    
    Args:
        summary (dict): High-level summary from the model.
        issues (list): List of issue dictionaries.
        
    Returns:
        Document: A python-docx Document object.
    """
    doc = Document()
    
    # Title
    title = doc.add_heading("Proofreading Summary Report", level=0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    
    # Overall Assessment
    if summary.get("overall_assessment"):
        doc.add_heading("Overall Assessment", level=1)
        doc.add_paragraph(summary["overall_assessment"])
    
    # Strengths
    if summary.get("strengths"):
        doc.add_heading("Strengths", level=1)
        for strength in summary["strengths"]:
            doc.add_paragraph(strength, style="List Bullet")
    
    # Areas for Improvement
    if summary.get("areas_for_improvement"):
        doc.add_heading("Areas for Improvement", level=1)
        for area in summary["areas_for_improvement"]:
            doc.add_paragraph(area, style="List Bullet")
    
    # Section-level Feedback
    if summary.get("sections"):
        doc.add_heading("Section Feedback", level=1)
        for section in summary["sections"]:
            name = section.get("name", "Untitled")
            feedback = section.get("feedback", "No specific feedback.")
            doc.add_heading(name, level=2)
            doc.add_paragraph(feedback)
    
    # Statistics
    if issues:
        doc.add_heading("Correction Statistics", level=1)
        total = len(issues)
        categories = [issue.get("category", "Other") for issue in issues]
        category_counts = Counter(categories)
        
        p = doc.add_paragraph()
        p.add_run(f"Total corrections applied: {total}").bold = True
        
        doc.add_paragraph("By category:")
        for cat, count in sorted(category_counts.items()):
            doc.add_paragraph(f"{cat}: {count}", style="List Bullet")
        
        doc.add_paragraph()
        note = doc.add_paragraph("See the tracked changes in the output document for detailed corrections.")
        note.runs[0].italic = True
    else:
        doc.add_heading("Corrections", level=1)
        doc.add_paragraph("No corrections needed. Great job!")
    
    return doc


def save_report(doc, output_path):
    """
    Save the report document to a file.
    
    Args:
        doc (Document): The python-docx Document object.
        output_path (str): Path to save the .docx file.
    """
    doc.save(output_path)


